"""Prepare read-only, Id-based UAT updates for Account_NEW commercial fields.

Writes customer data only under the operating system temporary directory.
The script never changes Salesforce records or calls Mexal.
"""

from __future__ import annotations

import csv
import hashlib
import json
import sys
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl

from prepare_account_new_import import ORG, SOURCE_SHEET, sf_json


FIELD_PATH = (
    Path(__file__).resolve().parents[1]
    / "force-app/main/default/objects/Account/fields/Tipologia_Attivita__c.field-meta.xml"
)
CATEGORY_VALUES = {"1", "4", "10", "14", "16", "17", "19"}


def value(row: dict, name: str) -> str:
    raw = row[name]
    return "" if raw is None else str(raw).strip()


def write_csv(path: Path, rows: list[dict], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("usage: prepare_account_commercial_update.py WORKBOOK.xlsx")
    workbook_path = Path(sys.argv[1]).resolve()
    if not workbook_path.is_file():
        raise RuntimeError("Workbook not found")

    org = sf_json("org", "display")
    if "sandbox.my.salesforce.com" not in org.get("instanceUrl", ""):
        raise RuntimeError("Refusing to prepare an update for a non-sandbox org")

    result = sf_json(
        "data",
        "query",
        "--query",
        "SELECT Id, External_CRM_ID__c, RecordType.DeveloperName "
        "FROM Account WHERE External_CRM_ID__c != null",
    )
    records = result["records"]
    if len(records) != result["totalSize"]:
        raise RuntimeError("Salesforce Account query was incomplete")
    org_by_crm = {record["External_CRM_ID__c"]: record for record in records}
    if len(org_by_crm) != len(records):
        raise RuntimeError("Duplicate CRM IDs in UAT")
    if any(record["RecordType"]["DeveloperName"] != "Azienda" for record in records):
        raise RuntimeError("An imported UAT Account is not an Azienda")

    namespace = {"m": "http://soap.sforce.com/2006/04/metadata"}
    field_xml = ET.parse(FIELD_PATH)
    allowed_activities = {
        element.text
        for element in field_xml.findall(".//m:value/m:fullName", namespace)
    }

    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    if SOURCE_SHEET not in workbook.sheetnames:
        raise RuntimeError("Account_NEW sheet is missing")
    iterator = workbook[SOURCE_SHEET].iter_rows(values_only=True)
    headers = next(iterator)
    required = {
        "CRM",
        "Partita Iva",
        "Categoria Provvigioni Cliente",
        "Tipologia Attività",
        "Zona",
    }
    if not required.issubset(headers):
        raise RuntimeError("Account_NEW headers changed")
    source = [dict(zip(headers, cells)) for cells in iterator]
    eligible = [row for row in source if value(row, "Partita Iva")]
    source_ids = [value(row, "CRM") for row in eligible]
    if len(source) != 8597 or len(eligible) != 8140:
        raise RuntimeError("Source row counts changed; review the extract")
    if len(set(source_ids)) != len(source_ids) or set(source_ids) != set(org_by_crm):
        raise RuntimeError("Source and UAT CRM ID sets differ; no update prepared")

    with_activity: list[dict] = []
    category_only: list[dict] = []
    blank_categories = 0
    for row in eligible:
        category = value(row, "Categoria Provvigioni Cliente")
        activity = value(row, "Tipologia Attività")
        if value(row, "Zona"):
            raise RuntimeError("Zona is no longer blank; review mapping before update")
        if category and category not in CATEGORY_VALUES:
            raise RuntimeError("Unexpected commission category")
        if activity and activity not in allowed_activities:
            raise RuntimeError("Activity value absent from field metadata")
        if not category:
            blank_categories += 1
            if activity:
                raise RuntimeError("Activity without a commission category")
            continue
        account_id = org_by_crm[value(row, "CRM")]["Id"]
        if activity:
            with_activity.append(
                {
                    "Id": account_id,
                    "Categoria_Provvigioni_Cliente__c": category,
                    "Tipologia_Attivita__c": activity,
                }
            )
        else:
            category_only.append(
                {"Id": account_id, "Categoria_Provvigioni_Cliente__c": category}
            )

    if len(with_activity) != 181 or len(category_only) != 7957 or blank_categories != 2:
        raise RuntimeError("Mapped field counts changed; review the extract")
    output = Path(tempfile.mkdtemp(prefix="pienissimo-commercial-update-"))
    write_csv(
        output / "sample.csv",
        with_activity[:1],
        ["Id", "Categoria_Provvigioni_Cliente__c", "Tipologia_Attivita__c"],
    )
    write_csv(
        output / "with_activity.csv",
        with_activity[1:],
        ["Id", "Categoria_Provvigioni_Cliente__c", "Tipologia_Attivita__c"],
    )
    write_csv(
        output / "category_only.csv",
        category_only,
        ["Id", "Categoria_Provvigioni_Cliente__c"],
    )
    manifest = {
        "org": ORG,
        "sheet": SOURCE_SHEET,
        "workbook_sha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
        "eligible_rows": len(eligible),
        "sample_rows": 1,
        "with_activity_after_sample": len(with_activity) - 1,
        "category_only_rows": len(category_only),
        "blank_category_rows": blank_categories,
        "mexal_trigger_fields_in_csv": [],
    }
    (output / "manifest.json").write_text(
        json.dumps(manifest, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps({"temporary_output_dir": str(output), **manifest}, indent=2))


if __name__ == "__main__":
    main()
