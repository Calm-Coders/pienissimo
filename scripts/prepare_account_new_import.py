"""Prepare the Pienissimo UAT Account import from Account_NEW.

Customer data is written only to a temporary directory. The script performs
read-only Salesforce queries and does not submit the import jobs.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl


ORG = "Pienissimo UAT"
SOURCE_SHEET = "Account_NEW"
NEW_FIELDS = [
    "Codice_Cliente_Mexal__c",
    "External_CRM_ID__c",
    "Name",
    "Partita_IVA__c",
    "Codice_Fiscale__c",
    "PEC__c",
    "Codice_Destinatario_SDI__c",
    "BillingStreet",
    "BillingPostalCode",
    "BillingCity",
    "BillingState",
    "BillingCountry",
    "ShippingStreet",
    "ShippingPostalCode",
    "ShippingCity",
    "ShippingState",
    "ShippingCountry",
    "Email__c",
    "Codice_ATECO__c",
    "Descrizione_ATECO__c",
    "Azienda_Test__c",
    "Codice_Agente_Esterno__c",
    "Agente__c",
    "RecordTypeId",
    "OwnerId",
]


def sf_json(*args: str) -> dict:
    sf_executable = shutil.which("sf.cmd") or shutil.which("sf")
    if not sf_executable:
        raise RuntimeError("Salesforce CLI was not found")
    result = subprocess.run(
        [sf_executable, *args, "--target-org", ORG, "--json"],
        capture_output=True,
        text=True,
        check=False,
    )
    try:
        payload = json.loads(result.stdout)
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"Salesforce CLI did not return JSON: {result.stderr[-300:]}") from exc
    if result.returncode or payload.get("status") != 0:
        raise RuntimeError(f"Salesforce CLI failed: {payload.get('message', payload.get('name'))}")
    return payload["result"]


def query(soql: str) -> list[dict]:
    result = sf_json("data", "query", "--query", soql)
    if len(result["records"]) != result["totalSize"]:
        raise RuntimeError("SOQL result was incomplete")
    return result["records"]


def value(row: dict, name: str) -> str:
    raw = row[name]
    return "" if raw is None else str(raw).strip()


def write_csv(path: Path, rows: list[dict], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("usage: prepare_account_new_import.py WORKBOOK.xlsx")
    workbook_path = Path(sys.argv[1]).resolve()
    if not workbook_path.is_file():
        raise RuntimeError("Workbook not found")

    org = sf_json("org", "display")
    if "sandbox.my.salesforce.com" not in org.get("instanceUrl", ""):
        raise RuntimeError("Refusing to prepare an import for a non-sandbox org")

    agents = query(
        "SELECT Id, Name, IsActive, Profile.Name FROM User "
        "WHERE Profile.Name = 'Agente'"
    )
    agent_by_name = {user["Name"]: user for user in agents}
    if len(agent_by_name) != len(agents) or any(user["IsActive"] for user in agents):
        raise RuntimeError("Agent users are ambiguous or active")

    owners = query(
        "SELECT Id, Name, IsActive FROM User "
        "WHERE Name = 'Amministratore Pienissimo'"
    )
    if len(owners) != 1 or not owners[0]["IsActive"]:
        raise RuntimeError("Active UAT Account owner is not unique")
    owner_id = owners[0]["Id"]

    types = query(
        "SELECT Id, DeveloperName FROM RecordType "
        "WHERE SObjectType = 'Account' AND DeveloperName = 'Azienda'"
    )
    if len(types) != 1:
        raise RuntimeError("Azienda Account record type is not unique")
    record_type_id = types[0]["Id"]

    existing = query(
        "SELECT Codice_Cliente_Mexal__c FROM Account "
        "WHERE Codice_Cliente_Mexal__c != null"
    )
    existing_keys = [account["Codice_Cliente_Mexal__c"] for account in existing]
    if len(existing_keys) != len(set(existing_keys)):
        raise RuntimeError("Duplicate Mexal external keys already exist in UAT")
    existing_key_set = set(existing_keys)

    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    if SOURCE_SHEET not in workbook.sheetnames:
        raise RuntimeError("Account_NEW sheet is missing")
    iterator = workbook[SOURCE_SHEET].iter_rows(values_only=True)
    headers = next(iterator)
    required = {
        "CRM",
        "Codice cliente esterno",
        "Ragione Sociale",
        "Partita Iva",
        "Azienda Proprietario Name",
        "Codice_agente",
        "Azienda Test",
    }
    if not required.issubset(headers):
        raise RuntimeError("Account_NEW headers changed")

    source_rows = [dict(zip(headers, cells)) for cells in iterator]
    new_rows: list[dict] = []
    existing_rows: list[dict] = []
    source_keys: set[str] = set()
    crm_ids: set[str] = set()
    filtered_without_vat = 0
    for source in source_rows:
        if not value(source, "Partita Iva"):
            filtered_without_vat += 1
            continue
        key = value(source, "Codice cliente esterno")
        crm_id = value(source, "CRM")
        name = value(source, "Ragione Sociale")
        agent_name = value(source, "Azienda Proprietario Name")
        if not key or not crm_id or not name or agent_name not in agent_by_name:
            raise RuntimeError("Required key, name, CRM ID, or agent user is missing")
        if key in source_keys or crm_id in crm_ids:
            raise RuntimeError("Duplicate Mexal key or CRM ID in filtered source")
        if not re.fullmatch(r"\d{18}", crm_id):
            raise RuntimeError("CRM ID is not an 18-digit text identifier")
        source_keys.add(key)
        crm_ids.add(crm_id)
        test_flag = value(source, "Azienda Test")
        if test_flag not in {"Yes", "No"}:
            raise RuntimeError("Unexpected Azienda Test value")

        account = {
            "Codice_Cliente_Mexal__c": key,
            "External_CRM_ID__c": crm_id,
            "Name": name,
            "Partita_IVA__c": value(source, "Partita Iva"),
            "Codice_Fiscale__c": value(source, "Codice Fiscale"),
            "PEC__c": value(source, "PEC"),
            "Codice_Destinatario_SDI__c": value(source, "SDI"),
            "BillingStreet": value(source, "Via fatturazione"),
            "BillingPostalCode": value(source, "Cap"),
            "BillingCity": value(source, "Città di fatturazione"),
            "BillingState": value(source, "Provincia di fatturazione"),
            "BillingCountry": value(source, "Paese di fatturazione"),
            "ShippingStreet": value(source, "Via spedizione"),
            "ShippingPostalCode": value(source, "CAP di spedizione"),
            "ShippingCity": value(source, "Città di spedizione"),
            "ShippingState": value(source, "Provincia di spedizione"),
            "ShippingCountry": value(source, "Paese di spedizione"),
            "Email__c": value(source, "E-mail Amministrativa"),
            "Codice_ATECO__c": value(source, "Codice Ateco"),
            "Descrizione_ATECO__c": value(source, "Ateco Desc"),
            "Azienda_Test__c": "true" if test_flag == "Yes" else "false",
            "Codice_Agente_Esterno__c": value(source, "Codice_agente"),
            "Agente__c": agent_by_name[agent_name]["Id"],
            "RecordTypeId": record_type_id,
            "OwnerId": owner_id,
        }
        if key in existing_key_set:
            existing_rows.append(account)
        else:
            new_rows.append(account)

    if len(source_rows) != 8597 or len(new_rows) + len(existing_rows) != 8140:
        raise RuntimeError("Source row counts changed; review the new extract")
    if len(agent_by_name) != 9:
        raise RuntimeError("Expected exactly nine inactive agent users")

    output = Path(tempfile.mkdtemp(prefix="pienissimo-account-import-"))
    write_csv(output / "sample.csv", new_rows[:1], NEW_FIELDS)
    write_csv(output / "new.csv", new_rows[1:], NEW_FIELDS)
    write_csv(output / "existing.csv", existing_rows, NEW_FIELDS[:-1])
    manifest = {
        "org": ORG,
        "sheet": SOURCE_SHEET,
        "workbook_sha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
        "source_rows": len(source_rows),
        "excluded_without_vat": filtered_without_vat,
        "eligible_rows": len(new_rows) + len(existing_rows),
        "sample_rows": 1,
        "new_rows_after_sample": len(new_rows) - 1,
        "existing_rows": len(existing_rows),
        "agent_users": len(agent_by_name),
    }
    (output / "manifest.json").write_text(
        json.dumps(manifest, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps({"temporary_output_dir": str(output), **manifest}, indent=2))


if __name__ == "__main__":
    main()
